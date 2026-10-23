"""Genera datos.js (lo que lee el mapa) a partir del Excel de puntos, y antes lo revisa.

Uso:  python herramientas/generar_datos.py [ruta_del_excel]
Por defecto lee  ../Puntos mapa bonos.xlsx  (carpeta mapa_bonos, fuera del repo).

Lee las hojas Puntos, Núcleos plano, Textos y Frases. No exporta las columnas internas
(nombre en el bono, precisión, fuente, comentarios).
Si encuentra ERRORES no escribe nada y termina con código 1. Los AVISOS no paran la generación.
"""
import json
import os
import re
import sys
from datetime import datetime

from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RAIZ, "..", "Puntos mapa bonos.xlsx")
IDIOMAS = ["es", "ca", "fr", "en"]
ESTACIONES = {"GV", "PA", "OA"}
SECTORES = {"GV", "NP"}  # GV = Grandvalira; NP = Nord Pass (Pal Arinsal + Ordino Arcalís)
TIPOS = {"Escuela": "escuela", "Restaurante": "restaurante", "Tienda": "tienda"}
ANDORRA = (42.40, 42.70, 1.38, 1.80)  # lat mín, lat máx, lng mín, lng máx

errores, avisos = [], []


def texto(x):
    return x.strip() if isinstance(x, str) else ("" if x is None else x)


def leer_hoja(wb, nombre):
    if nombre not in wb.sheetnames:
        errores.append(f"Falta la hoja '{nombre}' en el Excel.")
        return [], {}
    filas = list(wb[nombre].iter_rows(values_only=True))
    cab = [texto(c) for c in filas[0]]
    return filas[1:], {n: i for i, n in enumerate(cab) if n}


wb = load_workbook(EXCEL, data_only=True)

# ---------- Núcleos (posiciones sobre el plano) ----------
filas, col = leer_hoja(wb, "Núcleos plano")
nucleos = {}
for f in filas:
    cod = texto(f[0])
    if not cod or not isinstance(f[3], (int, float)):
        continue
    if cod in nucleos:
        errores.append(f"Núcleos plano: el código {cod} está repetido.")
    nucleos[cod] = {"est": texto(f[1]), "nombre": texto(f[2]), "x": f[3], "y": f[4]}

# ---------- Puntos ----------
filas, col = leer_hoja(wb, "Puntos")


def v(fila, nombre):
    return texto(fila[col[nombre]]) if nombre in col else ""


def si_no(fila, nombre, id_):
    x = v(fila, nombre)
    if x not in ("", "Sí", "No"):
        errores.append(f"{id_}: la columna '{nombre}' tiene '{x}'; tiene que ser Sí, No o estar vacía.")
    return x == "Sí"


puntos, ids, usados = [], set(), set()
for n_fila, f in enumerate(filas, start=2):
    id_ = v(f, "ID")
    if not id_:
        continue
    if id_ in ids:
        errores.append(f"Puntos fila {n_fila}: el ID {id_} está repetido.")
    ids.add(id_)
    tipo = TIPOS.get(v(f, "Tipo"))
    if not tipo:
        errores.append(f"{id_}: Tipo '{v(f, 'Tipo')}' no válido (Escuela, Restaurante o Tienda).")
        continue
    est = v(f, "Estación")
    if est not in ESTACIONES and not (tipo == "tienda" and est == "Todas"):
        errores.append(f"{id_}: Estación '{est}' no válida (GV, PA, OA; las tiendas, Todas).")
    p = {
        "id": id_, "tipo": tipo, "est": est,
        "zona": v(f, "Zona (agrupa en el mapa)"),
        "ubic": v(f, "Núcleo / ubicación"),
        "nombre": v(f, "Nombre en el mapa"),
        "marca": v(f, "Marca"),
        "express": si_no(f, "Contramarca Express", id_),
        "self": si_no(f, "Contramarca Self", id_),
        "local": v(f, "Tipo de local"),
        "nota": v(f, "Qué incluye / nota cliente"),
        "partner": si_no(f, "Partner GV", id_),
        "dir": v(f, "Dirección"),
        "nucleo": v(f, "Núcleo en plano"),
        "sector": v(f, "Sector (GV / NP)") or ("GV" if est == "GV" else "NP"),
    }
    if p["sector"] not in SECTORES:
        errores.append(f"{id_}: Sector '{p['sector']}' no válido: tiene que ser GV (Grandvalira) o NP (Nord Pass).")
    if tipo == "tienda" and not v(f, "Sector (GV / NP)"):
        errores.append(f"{id_}: a la tienda le falta el Sector (GV o NP).")
    if not p["nombre"]:
        errores.append(f"{id_}: falta el 'Nombre en el mapa'.")
    if not p["zona"]:
        errores.append(f"{id_}: falta la Zona.")
    lat, lng = f[col["Latitud"]], f[col["Longitud"]]
    if isinstance(lat, (int, float)) and isinstance(lng, (int, float)):
        if not (ANDORRA[0] <= lat <= ANDORRA[1] and ANDORRA[2] <= lng <= ANDORRA[3]):
            errores.append(f"{id_}: la posición {lat}, {lng} cae fuera de Andorra (¿latitud y longitud cambiadas?).")
        p["lat"], p["lng"] = round(lat, 6), round(lng, 6)
    elif lat not in (None, "") or lng not in (None, ""):
        errores.append(f"{id_}: Latitud/Longitud no son números ('{lat}', '{lng}'). Usa punto o coma decimal de Excel, sin texto.")
    if tipo in ("escuela", "restaurante"):
        if p["nucleo"]:
            n = nucleos.get(p["nucleo"])
            if not n:
                errores.append(f"{id_}: el núcleo '{p['nucleo']}' no existe en la hoja Núcleos plano.")
            elif n["est"] != est:
                errores.append(f"{id_}: el núcleo {p['nucleo']} es de {n['est']} pero el punto es de {est}.")
            usados.add(p["nucleo"])
        else:
            avisos.append(f"{id_} ({p['nombre']}): sin 'Núcleo en plano'; sale en la lista pero no en el plano.")
    if tipo == "restaurante" and p["self"] and est != "GV":
        avisos.append(f"{id_}: acepta Self pero el menú Self solo existe en Grandvalira.")
    if tipo == "tienda" and "lat" not in p:
        avisos.append(f"{id_} ({p['nombre']}): tienda sin posición; sale en la lista pero no en el mapa.")
    puntos.append(p)

for cod in nucleos:
    if cod not in usados:
        avisos.append(f"Núcleo {cod} ({nucleos[cod]['nombre']}) no lo usa ningún punto.")

# ---------- Textos ----------
filas, col = leer_hoja(wb, "Textos")
textos = {}
for f in filas:
    clave = texto(f[0])
    if not clave:
        continue
    vals = {l: texto(f[i + 1]) for i, l in enumerate(IDIOMAS)}
    if not vals["es"]:
        errores.append(f"Textos: la clave {clave} no tiene texto en ES.")
        continue
    marcas_es = sorted(re.findall(r"\{\w+\}", vals["es"]))
    for l in IDIOMAS[1:]:
        if not vals[l]:
            avisos.append(f"Textos: {clave} sin traducir a {l.upper()} (se mostrará en español).")
        elif sorted(re.findall(r"\{\w+\}", vals[l])) != marcas_es:
            errores.append(f"Textos: en {clave} ({l.upper()}) falta o sobra algo como {{n}}: tiene que llevar {marcas_es}.")
    textos[clave] = {l: t for l, t in vals.items() if t}

# ---------- Frases ----------
filas, col = leer_hoja(wb, "Frases")
frases = []
for f in filas:
    es = texto(f[0])
    if es:
        frases.append([es] + [texto(f[i]) or es for i in range(1, 4)])
conocidas = {fr[0] for fr in frases} | {"Bar", "Fast food", "Fast & Good", "Self-service"}  # iguales en los 4 idiomas
for campo, nombre in (("nota", "Qué incluye / nota cliente"), ("local", "Tipo de local")):
    for x in sorted({p[campo] for p in puntos if p[campo]} - conocidas):
        avisos.append(f"Frases: '{x}' ({nombre}) no tiene traducción; en CA/FR/EN saldrá en español.")

# ---------- Planos ----------
planos = {}
info_planos = os.path.join(RAIZ, "planos", "info.json")
if os.path.exists(info_planos):
    with open(info_planos, encoding="utf-8") as fh:
        planos = json.load(fh)
else:
    errores.append("No existe planos/info.json: ejecuta antes herramientas/generar_planos.py.")

# ---------- Informe ----------
print(f"Excel: {EXCEL}")
for est in sorted(ESTACIONES):
    e = [p for p in puntos if p["est"] == est]
    print(f"  {est}: {sum(p['tipo'] == 'escuela' for p in e)} escuelas, "
          f"{sum(p['tipo'] == 'restaurante' for p in e)} restaurantes "
          f"({sum(p['tipo'] == 'restaurante' and p['express'] for p in e)} Express, "
          f"{sum(p['tipo'] == 'restaurante' and p['self'] for p in e)} Self)")
print(f"  Tiendas: {sum(p['tipo'] == 'tienda' for p in puntos)} · Textos: {len(textos)} · Frases: {len(frases)}")
if avisos:
    print(f"\nAVISOS ({len(avisos)}), no impiden publicar:")
    for a in avisos:
        print("  -", a)
if errores:
    print(f"\nERRORES ({len(errores)}). Corrígelos en el Excel y vuelve a ejecutar. No se ha cambiado nada:")
    for e in errores:
        print("  -", e)
    sys.exit(1)

# ---------- Escritura ----------
datos = {"generado": datetime.now().strftime("%Y-%m-%d %H:%M"), "planos": planos, "nucleos": nucleos,
         "textos": textos, "frases": frases}
salida = os.path.join(RAIZ, "datos.js")
with open(salida, "w", encoding="utf-8", newline="\n") as fh:
    fh.write(f"// Generado por herramientas/generar_datos.py el {datos['generado']} desde 'Puntos mapa bonos.xlsx'. No editar a mano.\n")
    fh.write("window.MAPA_BONOS_DATOS = {\n")
    for k, val in datos.items():
        fh.write(f"{json.dumps(k)}: {json.dumps(val, ensure_ascii=False, separators=(',', ':'))},\n")
    fh.write('"puntos": [\n')
    fh.write(",\n".join(json.dumps(p, ensure_ascii=False, separators=(",", ":")) for p in puntos))
    fh.write("\n]};\n")
print(f"\nOK: {len(puntos)} puntos -> {salida}")
